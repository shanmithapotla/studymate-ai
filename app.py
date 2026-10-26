import io
import json
import os
from pathlib import Path

import streamlit as st
from google import genai

MODEL = "gemini-3.8-flash"
MAX_FILE_SIZE_MB = 20
MAX_TEXT_CHARS = 100_000
SUPPORTED_TYPES = ["pdf", "txt", "md", "docx", "pptx", "xlsx", "xlsm", "xls", "csv", "rtf", "png", "jpg", "jpeg", "webp", "tif", "tiff"]

st.set_page_config(page_title="StudyMate AI", page_icon="📚", layout="wide")


def get_demo_flashcards():
    return {"cards": [
        {"question": "What is a variable in Python?", "answer": "A name that refers to a value."},
        {"question": "What is a Python list?", "answer": "An ordered, mutable collection of items."},
        {"question": "What does a for loop do?", "answer": "It iterates over items in an iterable."},
        {"question": "What is a function?", "answer": "A reusable block of code that performs a task."},
        {"question": "What does len() return?", "answer": "The number of items in an object."},
    ]}


def get_demo_quiz():
    return {"questions": [
        {"question": "Which data type stores True or False?", "options": ["int", "str", "bool", "list"], "answer": "bool", "explanation": "The bool type represents True or False."},
        {"question": "Which symbol starts a Python single-line comment?", "options": ["//", "#", "/*", "$"], "answer": "#", "explanation": "Python uses # for a single-line comment."},
        {"question": "What does len([10, 20, 30]) return?", "options": ["2", "3", "10", "30"], "answer": "3", "explanation": "The list contains three items."},
        {"question": "Which keyword defines a function?", "options": ["func", "function", "def", "make"], "answer": "def", "explanation": "Python uses def to define a function."},
        {"question": "Which collection is immutable?", "options": ["list", "set", "dictionary", "tuple"], "answer": "tuple", "explanation": "A tuple cannot be modified after creation."},
    ]}


def extract_uploaded_text(uploaded_file):
    name = uploaded_file.name
    suffix = Path(name).suffix.lower().lstrip(".")
    data = uploaded_file.getvalue()
    if len(data) > MAX_FILE_SIZE_MB * 1024 * 1024:
        raise ValueError(f"{name} is larger than {MAX_FILE_SIZE_MB} MB.")
    if suffix in ("txt", "md"):
        return data.decode("utf-8", errors="replace")
    if suffix == "pdf":
        try:
            from pypdf import PdfReader
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages).strip()
        except Exception:
            text = ""
        if text:
            return text
        try:
            import fitz
            import pytesseract
            from PIL import Image
            doc = fitz.open(stream=data, filetype="pdf")
            return "\n".join(pytesseract.image_to_string(Image.open(io.BytesIO(page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes("png")))) for page in doc).strip()
        except Exception:
            raise ValueError("No selectable PDF text found. Scanned PDFs need PyMuPDF, Pillow, pytesseract, and the Tesseract system program.")
    if suffix == "rtf":
        from striprtf.striprtf import rtf_to_text
        return rtf_to_text(data.decode("utf-8", errors="replace"))
    if suffix == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if suffix == "pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    parts.append(shape.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if suffix in ("xlsx", "xlsm", "xls", "csv"):
        import pandas as pd
        if suffix == "csv":
            return pd.read_csv(io.BytesIO(data), dtype=str, keep_default_na=False).to_string(index=False)
        sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=str)
        return "\n\n".join(f"Sheet: {sheet}\n{frame.fillna('').to_string(index=False)}" for sheet, frame in sheets.items())
    if suffix in ("png", "jpg", "jpeg", "webp", "tif", "tiff"):
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(io.BytesIO(data))).strip()
    raise ValueError(f"Unsupported file type: .{suffix}")


def ask_gemini(prompt):
    api_key = os.getenv("GEMINI_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("GEMINI_API_KEY is missing. Set it in Terminal before starting Streamlit.")
    try:
        client = genai.Client(api_key=api_key)
        response = client.models.generate_content(model=MODEL, contents=prompt, config={"response_mime_type": "application/json"})
    except Exception as exc:
        msg = str(exc)
        low = msg.lower()
        if "429" in msg or "resource_exhausted" in low or "quota" in low:
            raise RuntimeError("Gemini API quota/rate limit reached. Use Demo mode for sample content, or retry after the quota resets. Check https://ai.dev/rate-limit.") from exc
        if "503" in msg or "unavailable" in low:
            raise RuntimeError("Gemini is temporarily unavailable. Please wait and try again.") from exc
        raise RuntimeError(f"Gemini API error: {msg}") from exc
    if not getattr(response, "text", None):
        raise RuntimeError("Gemini returned an empty response.")
    try:
        result = json.loads(response.text)
    except json.JSONDecodeError as exc:
        raise RuntimeError("Gemini returned invalid JSON. Please try again.") from exc
    if not isinstance(result, dict):
        raise RuntimeError("Gemini returned an unexpected response format.")
    return result


st.title("📚 StudyMate AI")
st.caption("Turn your notes into flashcards and quizzes, and plan your revision.")
with st.sidebar:
    st.header("⚙️ Settings")
    demo_mode = st.toggle("Demo mode (no API calls)", value=st.session_state.get("demo_mode", False), help="Uses sample Python study questions, not your uploaded notes.")
    st.session_state["demo_mode"] = demo_mode
    st.caption(f"AI model: {MODEL}")
    st.caption(f"Maximum file size: {MAX_FILE_SIZE_MB} MB per file")

notes = st.text_area("📝 Paste your study notes (optional)", height=150, placeholder="Paste lecture notes, definitions, formulas, or textbook excerpts...")
uploaded_files = st.file_uploader("📂 Upload one or more study files", type=SUPPORTED_TYPES, accept_multiple_files=True)
uploaded_texts = []
if uploaded_files:
    st.subheader("📄 Uploaded files")
    for uploaded_file in uploaded_files:
        try:
            extracted = extract_uploaded_text(uploaded_file)
            if not extracted.strip():
                st.warning(f"No readable text found in {uploaded_file.name}.")
                continue
            uploaded_texts.append(f"\n--- FILE: {uploaded_file.name} ---\n{extracted}")
            st.success(f"Read successfully: {uploaded_file.name}")
            with st.expander(f"Preview {uploaded_file.name}"):
                st.text(extracted[:5000])
                if len(extracted) > 5000:
                    st.caption("Preview shortened.")
        except Exception as exc:
            st.error(f"Could not read {uploaded_file.name}: {exc}")

file_text = "\n\n".join(uploaded_texts)
c1, c2 = st.columns(2)
with c1:
    flashcard_button = st.button("✨ Generate Flashcards", use_container_width=True)
with c2:
    quiz_button = st.button("🧠 Generate Quiz", use_container_width=True)

if flashcard_button or quiz_button:
    if demo_mode:
        if flashcard_button:
            st.session_state["flashcards"] = get_demo_flashcards()["cards"]
            st.session_state["flashcard_source"] = "demo"
        if quiz_button:
            st.session_state["quiz"] = get_demo_quiz()["questions"]
            st.session_state["quiz_source"] = "demo"
        st.info("Demo mode uses sample Python questions, not your uploaded notes.")
    else:
        study_parts = [part.strip() for part in (notes, file_text) if part and part.strip()]
        study_text = "\n\n".join(study_parts)[:MAX_TEXT_CHARS]
        if not study_text:
            st.warning("Paste notes or upload at least one readable file, or enable Demo mode.")
        else:
            try:
                if flashcard_button:
                    prompt = f'''Create exactly 5 useful study flashcards using only information supported by the source.
Return valid JSON only: {{"cards":[{{"question":"A clear question","answer":"A concise answer"}}]}}
SOURCE:
{study_text}'''
                    with st.spinner("Generating flashcards..."):
                        result = ask_gemini(prompt)
                    cards = result.get("cards", [])
                    if not isinstance(cards, list) or not cards:
                        raise RuntimeError("Gemini did not return valid flashcards.")
                    st.session_state["flashcards"] = cards[:10]
                    st.session_state["flashcard_source"] = "notes"
                if quiz_button:
                    prompt = f'''Create exactly 5 multiple-choice quiz questions based only on the source.
Each question must have four options and one correct answer matching an option exactly.
Return valid JSON only: {{"questions":[{{"question":"Question text","options":["A","B","C","D"],"answer":"Correct option text","explanation":"Short explanation"}}]}}
SOURCE:
{study_text}'''
                    with st.spinner("Generating quiz..."):
                        result = ask_gemini(prompt)
                    questions = result.get("questions", [])
                    if not isinstance(questions, list) or not questions:
                        raise RuntimeError("Gemini did not return valid quiz questions.")
                    st.session_state["quiz"] = questions[:10]
                    st.session_state["quiz_source"] = "notes"
            except Exception as exc:
                st.error(str(exc))

if st.session_state.get("flashcards"):
    st.divider()
    st.subheader("✨ Your Flashcards")
    if st.session_state.get("flashcard_source") == "demo":
        st.caption("Demo sample — not generated from your notes.")
    cards = st.session_state["flashcards"]
    if "flashcard_index" not in st.session_state:
        st.session_state["flashcard_index"] = 0
    idx = st.session_state["flashcard_index"] % len(cards)
    card = cards[idx]
    st.markdown(f"**Card {idx + 1} of {len(cards)}**")
    st.info(card.get("question", card.get("front", "Question unavailable")))
    with st.expander("Reveal answer"):
        st.write(card.get("answer", card.get("back", "Answer unavailable")))
    p, n, cl = st.columns(3)
    with p:
        if st.button("⬅️ Previous card"):
            st.session_state["flashcard_index"] = (idx - 1) % len(cards)
            st.rerun()
    with n:
        if st.button("Next card ➡️"):
            st.session_state["flashcard_index"] = (idx + 1) % len(cards)
            st.rerun()
    with cl:
        if st.button("Clear flashcards"):
            st.session_state.pop("flashcards", None)
            st.session_state.pop("flashcard_index", None)
            st.rerun()

if st.session_state.get("quiz"):
    st.divider()
    st.subheader("🧠 Practice Quiz")
    if st.session_state.get("quiz_source") == "demo":
        st.caption("Demo sample — Python questions, not generated from your notes.")
    questions = st.session_state["quiz"]
    with st.form("study_quiz_form"):
        answers = {}
        for i, q in enumerate(questions):
            st.markdown(f"**Q{i + 1}. {q.get('question', 'Question unavailable')}**")
            options = q.get("options", [])
            if not isinstance(options, list) or not options:
                options = ["No options provided"]
            answers[i] = st.radio("Choose your answer:", options, key=f"quiz_{i}_{st.session_state.get('quiz_source', 'notes')}", index=None)
        submitted = st.form_submit_button("Check answers", use_container_width=True)
    if submitted:
        score = sum(1 for i, q in enumerate(questions) if answers[i] is not None and str(answers[i]).strip().casefold() == str(q.get("answer", "")).strip().casefold())
        st.success(f"Your score: {score} / {len(questions)}")
        for i, q in enumerate(questions):
            correct = str(q.get("answer", "Not provided"))
            if answers[i] is not None and str(answers[i]).strip().casefold() == correct.casefold():
                st.markdown(f"**Q{i + 1}: Correct!**")
            else:
                st.markdown(f"**Q{i + 1}: Correct answer: {correct}**")
            if q.get("explanation"):
                st.caption(q["explanation"])

st.divider()
st.subheader("🗓️ Quick Study Planner")
st.caption("Tasks stay in this browser session; permanent database saving is not configured.")
with st.form("planner_form", clear_on_submit=True):
    task = st.text_input("Subject or task", placeholder="e.g., Revise DLD flip-flop conversions")
    due_date = st.date_input("Study date")
    priority = st.selectbox("Priority", ["High", "Medium", "Low"])
    add_task = st.form_submit_button("Add study task")
if "planner_tasks" not in st.session_state:
    st.session_state["planner_tasks"] = []
if add_task:
    if task.strip():
        st.session_state["planner_tasks"].append({"task": task.strip(), "date": due_date.isoformat(), "priority": priority, "done": False})
        st.success("Study task added.")
    else:
        st.warning("Enter a task first.")
if st.session_state["planner_tasks"]:
    st.markdown("**Your tasks**")
    for i, item in enumerate(st.session_state["planner_tasks"]):
        left, right = st.columns([5, 1])
        with left:
            st.write(f"{'✅' if item['done'] else '⬜'} **{item['task']}** — {item['date']} · {item['priority']}")
        with right:
            if not item["done"] and st.button("Done", key=f"done_{i}"):
                item["done"] = True
                st.rerun()
    if st.button("Clear all planner tasks"):
        st.session_state["planner_tasks"] = []
        st.rerun()

st.caption("StudyMate AI • Verify generated material against your course notes.")
